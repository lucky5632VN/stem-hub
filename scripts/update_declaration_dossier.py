import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_document_defaults(doc):
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    
    p_format = style.paragraph_format
    p_format.line_spacing = 1.2
    p_format.space_after = Pt(6)

def add_title(doc, title_text, subtitle_text="", date_text=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title_text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(17)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x00, 0x28, 0x55)
    
    if subtitle_text:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(4)
        run2 = p2.add_run(subtitle_text)
        run2.font.name = 'Times New Roman'
        run2.font.size = Pt(13)
        run2.font.bold = True
        run2.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    if date_text:
        p3 = doc.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_after = Pt(14)
        run3 = p3.add_run(date_text)
        run3.font.name = 'Times New Roman'
        run3.font.size = Pt(11)
        run3.font.italic = True
        run3.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x00, 0x28, 0x55)
    return p

def add_bullet(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r1 = p.add_run(bold_prefix + ": ")
        r1.font.name = 'Times New Roman'
        r1.font.bold = True
        r1.font.size = Pt(12)
        r1.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    r2 = p.add_run(text)
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(12)
    return p

def generate_declaration_docx(output_path):
    doc = Document()
    set_document_defaults(doc)
    
    add_title(
        doc,
        "KÊ KHAI CÔNG CỤ AI, DỮ LIỆU, API, THƯ VIỆN VÀ MÃ NGUỒN MỞ",
        "DỰ ÁN: STEM HUB — HỆ SINH THÁI HỌC TẬP VÀ PHÒNG LAB ẢO TÍCH HỢP AI",
        "Đơn vị: Trường THPT Quỳnh Lưu 4 — Tháng 09/2026"
    )
    
    p_preamble = doc.add_paragraph()
    p_preamble.paragraph_format.space_after = Pt(10)
    r_pre = p_preamble.add_run(
        "Văn bản kê khai trung thực, công khai 100% toàn bộ các công cụ Trí tuệ Nhân tạo (AI Tools), nguồn dữ liệu huấn luyện, "
        "giao diện lập trình ứng dụng (APIs), thư viện phần mềm và mã nguồn mở sử dụng trong dự án theo đúng thể lệ "
        "Cuộc thi Sáng tạo trẻ trong lĩnh vực Trí tuệ Nhân tạo do Ban Bí thư Trung ương Đoàn phát động."
    )
    r_pre.font.italic = True
    r_pre.font.size = Pt(11.5)
    r_pre.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    
    # I. CAM KẾT MINH BẠCH
    add_heading_1(doc, "I. CAM KẾT MINH BẠCH VÀ TRÁCH NHIỆM NGUỒN MỞ")
    p1 = doc.add_paragraph()
    p1.paragraph_format.space_after = Pt(6)
    p1.add_run(
        "Nhóm phát triển dự án STEM HUB (Trường THPT Quỳnh Lưu 4) xin cam kết công khai, minh bạch toàn bộ các công cụ AI, "
        "bộ dữ liệu, giao diện API và thư viện mã nguồn mở được tích hợp trong sản phẩm. Toàn bộ mã nguồn giao diện Sci-Fi HUD, "
        "thuật toán mạng nơ-ron 3 Tầng AI tự học, thuật toán di truyền Bot PvP và bộ ngân hàng dữ liệu câu hỏi 4 môn học "
        "hoàn toàn là sản phẩm sáng tạo chính chủ do nhóm tự thiết kế, lập trình và thẩm định khoa học."
    )
    
    # II. CÔNG CỤ AI
    add_heading_1(doc, "II. KÊ KHAI CÁC CÔNG CỤ TRÍ TUỆ NHÂN TẠO ĐÃ SỬ DỤNG (AI TOOLS)")
    
    table_ai = doc.add_table(rows=1, cols=4)
    table_ai.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_ai)
    
    hdr_cells = table_ai.rows[0].cells
    hdr_titles = ["STT", "Tên công cụ AI", "Đơn vị phát hành", "Mục đích & Phạm vi ứng dụng trong dự án"]
    col_widths = [Inches(0.5), Inches(1.8), Inches(1.5), Inches(2.7)]
    
    for i, title in enumerate(hdr_titles):
        cell = hdr_cells[i]
        cell.width = col_widths[i]
        set_cell_background(cell, "002855")
        set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i < 3 else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title)
        run.font.name = 'Times New Roman'
        run.font.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    ai_data = [
        ("1", "Antigravity AI Agent", "Google DeepMind", "Hỗ trợ nhóm rã bài toán kỹ thuật, viết mã nguồn (Next.js 16, React 19, TF.js), gỡ lỗi runtime, thiết kế giao diện Sci-Fi HUD và xây dựng bộ test audit cô lập."),
        ("2", "Google Gemini API\n(Flash 1.5 & Pro 1.5)", "Google AI", "Xử lý ngôn ngữ tự nhiên (Tầng 1 LLM), tư vấn kiến thức SGK THPT 2025 cho Trợ lý AI Nova và phân tích ngữ cảnh bài tập."),
        ("3", "Meta AudioCraft / MusicGen", "Meta AI", "Tạo nhạc nền học tập tập trung và sinh hiệu ứng âm thanh khoa học (SFX) cho các tương tác thí nghiệm và Đấu trường PvP."),
        ("4", "TensorFlow.js & Face-API", "Google / Open Source", "Nhận diện 68 điểm mốc khuôn mặt Cyan Neon (ResNet-34) xác thực 2FA sinh học và huấn luyện mô hình Online Learning Tầng 3 trực tiếp ở client-side."),
        ("5", "Transformers.js\n(@huggingface/transformers)", "Hugging Face", "Nghiên cứu và đóng gói pipeline mô hình ngôn ngữ thu nhỏ (SLM On-device) chạy qua WebGPU phục vụ mở rộng Offline.")
    ]
    
    for row_idx, data in enumerate(ai_data):
        row = table_ai.add_row()
        bg_color = "F8F9FA" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            cell.width = col_widths[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx < 3 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            if col_idx == 1:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # III. NGUỒN DỮ LIỆU
    add_heading_1(doc, "III. KÊ KHAI NGUỒN DỮ LIỆU HUẤN LUYỆN VÀ TRA CỨU")
    add_bullet(doc, "1. Bộ Dữ Liệu Tri Thức STEM THPT (303 đoạn tri thức / 9.061 token)", "Dùng cho Tầng 2 RAG Vector Memory tra cứu kiến thức 4 môn Toán, Lý, Hóa, Sinh SGK GDPT 2018 (Lớp 10-12). Nguồn gốc: Nhóm học sinh tự tổng hợp, chuẩn hóa và biên soạn chuẩn xác.")
    add_bullet(doc, "2. Ngân Hàng Câu Hỏi Trắc Nghiệm THPT 2025 (500+ câu hỏi)", "Phục vụ Đấu trường PvP 1v1 và phân hệ luyện tập, phân hóa độ khó IRT Rasch theo điểm ELO. Nguồn gốc: Nhóm tự biên soạn, kiểm duyệt đáp án theo ma trận đề thi tốt nghiệp mới của Bộ GD&ĐT.")
    add_bullet(doc, "3. Bộ Dữ Liệu Phản Ứng Hóa Học (144 mẫu phản ứng ChemLab)", "Dùng cho Canvas 2D ChemLab phân tích chất hóa học, mô phỏng màu sắc kết tủa, chỉ thị màu và phát cảnh báo an toàn. Nguồn gốc: Nhóm tự tổng hợp từ Bảng tính tan và Sổ tay Hóa học THPT.")
    add_bullet(doc, "4. Bộ Dữ Liệu Mẫu Ý Định Hội Thoại (392 câu thoại Tiếng Việt)", "Huấn luyện mô hình Naive Bayes Intent Classifier cho Trợ lý Nova nhận diện giọng nói, câu hỏi tắt và teencode của học sinh. Nguồn gốc: Nhóm tự thu thập thử nghiệm từ học sinh THPT Quỳnh Lưu 4.")
    add_bullet(doc, "5. Bộ Dữ Liệu Quần Thể 20 AI Bot PvP Test", "20 tài khoản Bot có điểm ELO từ 50 - 130 tiến hóa di truyền (Genetic Algorithm) thích ứng theo trình độ người chơi qua các thế hệ mà không bị reset. Nguồn gốc: Tự động tạo và lưu trữ trên Firestore & IndexedDB.")

    # IV. GIAO DIỆN API VÀ DỊCH VỤ TÍCH HỢP
    add_heading_1(doc, "IV. KÊ KHAI GIAO DIỆN API VÀ DỊCH VỤ TÍCH HỢP")
    add_bullet(doc, "1. Google Gemini REST API (@google/generative-ai)", "Giao tiếp với Gemini LLM để tổng hợp câu trả lời thân thiện, mạch lạc cho học sinh.")
    add_bullet(doc, "2. Cloudflare Workers API (Edge Reverse Proxy)", "Cổng trung gian Edge Computing chuyển tiếp yêu cầu gọi AI, bảo vệ Secret API Key không bị lộ qua Network tab và điều phối hạn ngạch 100.000 requests/ngày miễn phí.")
    add_bullet(doc, "3. Firebase Auth & Cloud Firestore (Google Firebase)", "Quản lý tài khoản học sinh, lưu trữ ai_metrics, model_snapshot và phòng đấu PvP.")
    add_bullet(doc, "4. HTML5 Web Worker API", "Khởi tạo luồng nền độc lập (Background Thread) tính toán ma trận RAG TF-IDF Cosine Similarity, giải phóng hoàn toàn Main Thread để duy trì giao diện mượt mà 60 FPS.")
    add_bullet(doc, "5. Web Speech API (SpeechRecognition)", "Tích hợp sẵn trên trình duyệt, nhận diện giọng nói Tiếng Việt chuyển thành văn bản (STT).")
    add_bullet(doc, "6. Web Audio API", "Phát âm thanh hiệu ứng thí nghiệm, xử lý Audio Ducking (tự giảm nhạc nền khi AI phát biểu).")
    add_bullet(doc, "7. IndexedDB Storage API", "Lưu trữ cục bộ ma trận trọng số TensorFlow.js Online Model và bộ đệm offline cache Firestore (persistentSingleTabManager), ngăn chặn lỗi tràn hạn mức LocalStorage 5MB.")

    # V. THƯ VIỆN VÀ MÃ NGUỒN MỞ
    add_heading_1(doc, "V. KÊ KHAI THƯ VIỆN VÀ MÃ NGUỒN MỞ")
    
    table_lib = doc.add_table(rows=1, cols=5)
    table_lib.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_lib)
    
    hdr_lib_cells = table_lib.rows[0].cells
    hdr_lib_titles = ["STT", "Thư viện / Framework", "Phiên bản", "Giấy phép (License)", "Mục đích sử dụng"]
    col_lib_widths = [Inches(0.4), Inches(1.8), Inches(0.9), Inches(1.2), Inches(2.2)]
    
    for i, title in enumerate(hdr_lib_titles):
        cell = hdr_lib_cells[i]
        cell.width = col_lib_widths[i]
        set_cell_background(cell, "002855")
        set_cell_margins(cell, top=120, bottom=120, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i < 4 else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title)
        run.font.name = 'Times New Roman'
        run.font.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    lib_data = [
        ("1", "Next.js", "16.2.6", "MIT", "Framework chính dựng Web Application & Serverless API Routes."),
        ("2", "React / React-DOM", "19.2.4", "MIT", "Thư viện xây dựng giao diện người dùng tương tác."),
        ("3", "TensorFlow.js (@tensorflow/tfjs)", "4.22.0", "Apache 2.0", "Mô hình Neural Network 3 Tầng & Xác thực Sinh học 2FA."),
        ("4", "Three.js / @react-three/fiber", "0.184.0 / 9.6.1", "MIT", "Mô phỏng 3D phòng thí nghiệm Vật lý quang học laser & con lắc."),
        ("5", "KaTeX", "0.18.4", "MIT", "Hiển thị công thức Toán, Lý, Hóa chuẩn LaTeX / SGK GDPT 2018."),
        ("6", "Tailwind CSS & Framer Motion", "4.0 / 12.39.0", "MIT", "Design System Sci-Fi HUD & hiệu ứng chuyển động mượt mà."),
        ("7", "Zustand", "5.0.13", "MIT", "Quản lý state toàn cục và State Batching cho Đấu trường PvP."),
        ("8", "Lucide React", "1.16.0", "ISC", "Bộ biểu tượng icon khoa học & giáo dục."),
        ("9", "html2canvas-pro", "2.0.4", "MIT", "Kết xuất ảnh chụp màn hình chứng nhận học tập và bảng điểm PvP."),
        ("10", "XLSX (SheetJS)", "0.18.5", "Apache 2.0", "Xử lý xuất nhập dữ liệu ngân hàng câu hỏi đề thi.")
    ]
    
    for row_idx, data in enumerate(lib_data):
        row = table_lib.add_row()
        bg_color = "F8F9FA" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            cell.width = col_lib_widths[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx < 4 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10.5)
            if col_idx == 1:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # VI. CAM KẾT BẢN QUYỀN VÀ TRÁCH NHIỆM PHÁP LÝ
    add_heading_1(doc, "VI. CAM KẾT BẢN QUYỀN VÀ TRÁCH NHIỆM PHÁP LÝ")
    p_commit = doc.add_paragraph()
    p_commit.paragraph_format.space_after = Pt(6)
    p_commit.add_run(
        "Nhóm phát triển dự án STEM HUB (Trường THPT Quỳnh Lưu 4) xin cam đoan trước Ban Bí thư Trung ương Đoàn và Ban Giám khảo Cuộc thi:\n"
    )
    add_bullet(doc, "1. Tuân thủ giấy phép bản quyền", "Toàn bộ các thư viện, framework mã nguồn mở nêu trên đều được sử dụng hợp pháp theo đúng điều khoản giấy phép mã nguồn mở tương ứng (MIT, Apache 2.0, ISC).")
    add_bullet(doc, "2. Tính chính chủ và nguyên bản", "Toàn bộ mã nguồn giao diện Sci-Fi HUD, thuật toán mạng nơ-ron 3 Tầng AI tự học, thuật toán di truyền Bot PvP và ngân hàng 500+ câu hỏi là sản phẩm sáng tạo chính chủ của nhóm học sinh Trường THPT Quỳnh Lưu 4.")
    add_bullet(doc, "3. Trách nhiệm pháp lý và đạo đức AI", "Nhóm tuyệt đối không sao chép nguyên mẫu, không vi phạm quyền sở hữu trí tuệ của bất kỳ tổ chức hay cá nhân nào, và hoàn toàn chịu trách nhiệm trước pháp luật về tính trung thực của văn bản kê khai này.")

    # VII. CHỮ KÝ XÁC NHẬN
    p_date = doc.add_paragraph()
    p_date.paragraph_format.space_before = Pt(14)
    p_date.paragraph_format.space_after = Pt(10)
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_d = p_date.add_run("Nghệ An, ngày 03 tháng 09 năm 2026")
    r_d.font.italic = True
    r_d.font.size = Pt(12)
    
    table_sig = doc.add_table(rows=1, cols=2)
    table_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    cell_gv = table_sig.rows[0].cells[0]
    cell_gv.width = Inches(3.2)
    p_gv = cell_gv.paragraphs[0]
    p_gv.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_gv1 = p_gv.add_run("GIÁO VIÊN HƯỚNG DẪN\n")
    r_gv1.bold = True
    r_gv1.font.size = Pt(12)
    r_gv1.font.color.rgb = RGBColor(0x00, 0x28, 0x55)
    r_gv2 = p_gv.add_run("(Ký và ghi rõ họ tên)\n\n\n\n\n")
    r_gv2.italic = True
    r_gv2.font.size = Pt(11)
    r_gv3 = p_gv.add_run("Thầy Nguyễn Văn Phúc")
    r_gv3.bold = True
    r_gv3.font.size = Pt(12)
    
    cell_team = table_sig.rows[0].cells[1]
    cell_team.width = Inches(3.2)
    p_team = cell_team.paragraphs[0]
    p_team.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t1 = p_team.add_run("ĐẠI DIỆN NHÓM TÁC GIẢ / ĐỘI TRƯỞNG\n")
    r_t1.bold = True
    r_t1.font.size = Pt(12)
    r_t1.font.color.rgb = RGBColor(0x00, 0x28, 0x55)
    r_t2 = p_team.add_run("(Ký và ghi rõ họ tên)\n\n\n\n\n")
    r_t2.italic = True
    r_t2.font.size = Pt(11)
    r_t3 = p_team.add_run("Hồ Hoàng Anh\n")
    r_t3.bold = True
    r_t3.font.size = Pt(12)
    r_t4 = p_team.add_run("Thành viên: Nguyễn Trung Hiếu, Đậu Việt Hưng")
    r_t4.font.size = Pt(11)
    r_t4.font.italic = True

    doc.save(output_path)
    print(f"✅ Đã lưu thành công văn bản kê khai: {output_path}")

def fix_javit_in_ho_so():
    doc_path = "e:/project/cuoc_thi_AI/AI2026_Mẫu hồ sơ.docx"
    doc = docx.Document(doc_path)
    
    modified = False
    for i, p in enumerate(doc.paragraphs):
        if "Javit" in p.text or "javit" in p.text:
            new_text = p.text.replace("Javit HUD", "NOVA HUD").replace("javit HUD", "NOVA HUD").replace("Javit", "NOVA")
            print(f"Sửa 'Javit' tại paragraph [{i}]:")
            print("  Cũ:", p.text[:80])
            print("  Mới:", new_text[:80])
            p.text = new_text
            modified = True
            
    if modified:
        doc.save(doc_path)
        print("✅ Đã chuẩn hóa toàn bộ 'Javit' -> 'NOVA' trong AI2026_Mẫu hồ sơ.docx!")
    else:
        print("Không còn chữ Javit nào trong AI2026_Mẫu hồ sơ.docx.")

if __name__ == "__main__":
    path1 = "e:/project/cuoc_thi_AI/Kê khai công cụ AI, dữ liệu, API, thư viện, mã nguồn mở.docx"
    path2 = "e:/project/cuoc_thi_AI/Ke_Khai_Cong_Cu_AI_Du_Lieu_API_Thu_Vien_Ma_Nguon_Mo.docx"
    path3 = "e:/project/cuoc_thi_AI/4_Ke_Khai_Cong_Cu_AI_Du_Lieu_API_Thu_Vien.docx"
    
    generate_declaration_docx(path1)
    generate_declaration_docx(path2)
    generate_declaration_docx(path3)
    
    fix_javit_in_ho_so()
    print("HOÀN TẤT NÂNG CẤP VĂN BẢN KÊ KHAI VÀ ĐỒNG BỘ HỒ SƠ!")
