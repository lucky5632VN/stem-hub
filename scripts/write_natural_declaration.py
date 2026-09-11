import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

def set_document_defaults(doc):
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    
    p_format = style.paragraph_format
    p_format.line_spacing = 1.2
    p_format.space_after = Pt(6)

def add_title(doc, title_text, subtitle_text="", school_text=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title_text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x00, 0x28, 0x55)
    
    if subtitle_text:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(3)
        run2 = p2.add_run(subtitle_text)
        run2.font.name = 'Times New Roman'
        run2.font.size = Pt(13)
        run2.font.bold = True
        run2.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    if school_text:
        p3 = doc.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_after = Pt(14)
        run3 = p3.add_run(school_text)
        run3.font.name = 'Times New Roman'
        run3.font.size = Pt(11.5)
        run3.font.italic = True
        run3.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x00, 0x28, 0x55)
    return p

def add_numbered_item(doc, num_prefix, bold_title, content):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.2
    
    r_num = p.add_run(f"{num_prefix}. ")
    r_num.font.name = 'Times New Roman'
    r_num.font.bold = True
    r_num.font.size = Pt(12)
    r_num.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    r_title = p.add_run(f"{bold_title}: ")
    r_title.font.name = 'Times New Roman'
    r_title.font.bold = True
    r_title.font.size = Pt(12)
    r_title.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    r_content = p.add_run(content)
    r_content.font.name = 'Times New Roman'
    r_content.font.size = Pt(12)
    return p

def add_paragraph(doc, text, italic=False, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.2
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.italic = italic
    run.font.bold = bold
    return p

def create_natural_student_doc(output_path):
    doc = Document()
    set_document_defaults(doc)
    
    add_title(
        doc,
        "BẢN KÊ KHAI CÔNG CỤ AI, NGUỒN DỮ LIỆU, API VÀ MÃ NGUỒN MỞ",
        "DỰ ÁN: STEM HUB — HỆ SINH THÁI HỌC TẬP VÀ PHÒNG THÍ NGHIỆM ẢO 3D",
        "Đơn vị: Nhóm học sinh Trường THPT Quỳnh Lưu 4, tỉnh Nghệ An"
    )
    
    # I. LỜI CAM ĐOAN MỞ ĐẦU
    add_heading(doc, "I. LỜI CAM ĐOAN MỞ ĐẦU CỦA NHÓM TÁC GIẢ")
    
    add_paragraph(
        doc,
        "Em là Hồ Hoàng Anh, học sinh lớp 12A trường THPT Quỳnh Lưu 4, đại diện cho nhóm tác giả gồm ba thành viên chúng em "
        "(Hồ Hoàng Anh, Nguyễn Trung Hiếu, Đậu Việt Hưng) dưới sự hướng dẫn của thầy Nguyễn Văn Phúc, xin được gửi lời chào trân trọng "
        "đến Ban tổ chức và Ban Giám khảo Cuộc thi Sáng tạo trẻ trong lĩnh vực Trí tuệ Nhân tạo năm 2026 do Ban Bí thư Trung ương Đoàn phát động."
    )
    add_paragraph(
        doc,
        "Nhóm chúng em xin cam đoan kê khai một cách đầy đủ, rõ ràng và trung thực nhất toàn bộ các công cụ trí tuệ nhân tạo, "
        "các nguồn dữ liệu học tập, các giao diện kết nối API và các thư viện mã nguồn mở mà nhóm đã tìm hiểu, tích hợp và sử dụng "
        "trong suốt quá trình xây dựng sản phẩm STEM HUB. Mọi tài nguyên công nghệ được sử dụng đều phục vụ mục đích học tập phi lợi nhuận "
        "cho học sinh phổ thông và hoàn toàn tuân thủ thể lệ của cuộc thi."
    )
    
    # II. CÔNG CỤ AI
    add_heading(doc, "II. CÁC CÔNG CỤ TRÍ TUỆ NHÂN TẠO NHÓM EM ĐÃ SỬ DỤNG")
    add_paragraph(
        doc,
        "Trong quá trình thực hiện dự án, nhóm chúng em đã chủ động tìm tòi và ứng dụng một số công cụ AI hiện đại để hỗ trợ hiện thực hóa các ý tưởng của mình:"
    )
    
    add_numbered_item(
        doc, "1", "Trợ lý lập trình Antigravity AI của Google",
        "Đây là công cụ AI hỗ trợ lập trình chính của nhóm. Chúng em đưa ra các yêu cầu về giao diện, thuật toán và bài toán thực tế của học sinh THPT, sau đó Antigravity AI giúp nhóm viết các đoạn mã nguồn bằng Next.js, React và TensorFlow.js. Công cụ này cũng giúp chúng em tìm và sửa các lỗi cú pháp, tối ưu các đoạn code để trang web chạy mượt mà hơn. Toàn bộ ý tưởng bài toán, kịch bản giao diện và việc kiểm tra chất lượng sản phẩm đều do nhóm chúng em trực tiếp thực hiện và làm chủ."
    )
    add_numbered_item(
        doc, "2", "Mô hình ngôn ngữ Google Gemini API",
        "Nhóm em tích hợp mô hình Gemini (bản Flash 1.5 và Pro 1.5) để vận hành Trợ lý ảo Nova. Khi các bạn học sinh vào trang web hỏi về kiến thức các môn Toán, Lý, Hóa, Sinh hoặc nhờ hướng dẫn giải bài tập, Gemini sẽ đóng vai trò như một người gia sư thân thiện, giải thích các hiện tượng khoa học theo đúng tinh thần sách giáo khoa mới và gợi ý các bạn mở đúng phòng thí nghiệm tương ứng."
    )
    add_numbered_item(
        doc, "3", "Mô hình Học máy Tự huấn luyện cục bộ (Self-Trained AI) và Bảo mật sinh trắc học",
        "Thay vì chỉ gọi các API có sẵn từ bên ngoài, nhóm chúng em đã chủ động nghiên cứu và tự lập trình (tự code) 6 mô hình Trí tuệ Nhân tạo chuyên biệt chạy 100% trên trình duyệt phía người dùng (Client-Side AI Execution) bằng JavaScript thuần và thư viện TensorFlow.js, hoàn toàn không tốn chi phí thuê máy chủ GPU:\n"
        "• Mạng Nơ-ron Phân loại Đa lớp Softmax (Softmax Multi-Class Neural Network): Được nhóm tự huấn luyện cho Phòng Thí nghiệm Hóa học ảo, có nhiệm vụ phân tích các thông số nhiệt độ, áp suất, độ pH và nồng độ chất để dự báo sớm mức độ an toàn và các rủi ro hóa chất (cháy nổ, tỏa nhiệt mạnh, sinh khí độc).\n"
        "• Mô hình k-Láng giềng gần nhất (k-NN - k-Nearest Neighbors): Được tích hợp trong Phòng Thí nghiệm Vật lý ảo nhằm đối chiếu số liệu đo đạc thực tế của học sinh (như chu kỳ dao động con lắc, góc khúc xạ ánh sáng) với các phân cụm dữ liệu chuẩn, từ đó tự động chẩn đoán các dạng sai số thực nghiệm phổ biến.\n"
        "• Bộ phân loại Ý định Đa thức Naive Bayes (Multinomial Naive Bayes Intent Classifier): Hoạt động như \"bộ não\" điều phối của Trợ lý ảo Nova, giúp phân tích xác suất từ khóa trong câu hỏi của học sinh để tự động phân luồng chính xác ý định (hỏi lý thuyết, mở thí nghiệm, tính toán hay tra cứu).\n"
        "• Động cơ Đánh giá Thích ứng Năng lực 1PL Rasch IRT (Item Response Theory): Triển khai mô hình trắc lượng tri thức cổ điển 1 tham số trong phân hệ Toán học, liên tục ước lượng năng lực tiềm ẩn (θ) của học sinh dựa trên ma trận đúng/sai để tự động đề xuất câu hỏi có độ khó (b) tối ưu nhất.\n"
        "• Mạng Suy luận Xác suất Bayes với Bảng Xác suất Có điều kiện (Bayesian Belief Network & CPT - Conditional Probability Table): Xây dựng trong phân hệ Di truyền học (Sinh học) để mô hình hóa sơ đồ phả hệ gia đình, tự động tính toán xác suất di truyền bệnh và tỷ lệ phân ly kiểu gen, kiểu hình theo đúng quy luật Menđen.\n"
        "• Động cơ Trích xuất Tri thức Tăng cường Cục bộ (In-Browser RAG Engine): Thực hiện tìm kiếm và chọn lọc ngữ cảnh bài học phù hợp trong bộ nhớ tạm trước khi gửi cho trợ lý ảo sinh phản hồi.\n"
        "Bên cạnh 6 mô hình tự huấn luyện trên, nhóm còn tích hợp mạng nơ-ron tích chập ResNet-34 thông qua thư viện Face-API để nhận diện khuôn mặt học sinh phục vụ điểm danh và bảo mật hai lớp (Biometric 2FA). Toàn bộ 128 đặc trưng vector khuôn mặt đều được trích xuất và so khớp trực tiếp trên trình duyệt qua webcam, tuyệt đối không gửi hình ảnh khuôn mặt của học sinh lên mạng Internet."
    )
    add_numbered_item(
        doc, "4", "Công cụ âm thanh Meta AudioCraft",
        "Nhóm chúng em đã sử dụng công cụ AI này để tạo ra một số bản nhạc nền nhẹ nhàng giúp học sinh tập trung khi ngồi học bài, cùng các âm thanh mô phỏng tiếng bọt khí khi hóa chất tác dụng với nhau hoặc tiếng laser quang học để các phòng thí nghiệm ảo nghe sống động và chân thực hơn."
    )
    add_numbered_item(
        doc, "5", "Thư viện Transformers.js của Hugging Face",
        "Nhóm em đưa thư viện này vào mã nguồn dự án để nghiên cứu việc đóng gói các mô hình ngôn ngữ thu nhỏ chạy trực tiếp trên card đồ họa của máy tính học sinh qua công nghệ WebGPU, phục vụ cho hướng phát triển chạy hoàn toàn ngoại tuyến trong tương lai."
    )
    
    # III. NGUỒN DỮ LIỆU
    add_heading(doc, "III. NGUỒN DỮ LIỆU NHÓM EM TỰ THU THẬP VÀ BIÊN SOẠN")
    add_paragraph(
        doc,
        "Để hệ thống có đầy đủ kiến thức chính xác phục vụ học tập, nhóm chúng em đã tự tay xây dựng và chuẩn hóa các bộ dữ liệu sau:"
    )
    
    add_numbered_item(
        doc, "1", "Cơ sở tri thức 4 môn học THPT (303 đoạn tri thức)",
        "Nhóm em đã đọc kỹ sách giáo khoa lớp 10, 11, 12 theo chương trình giáo dục phổ thông 2018 và các tài liệu ôn thi của Bộ Giáo dục và Đào tạo, sau đó tự tóm tắt thành 303 đoạn tri thức trọng tâm cho 4 môn Toán, Lý, Hóa, Sinh (khoảng hơn 9.000 từ khóa). Đặc biệt, nhóm đã lập trình thuật toán Ma trận Tần suất Từ - Nghịch đảo Tần suất Văn bản (TF-IDF - Term Frequency-Inverse Document Frequency) kết hợp cùng hàm đo khoảng cách Độ tương đồng Cosine (Cosine Similarity) để trích xuất ngữ cảnh RAG trực tiếp trên trình duyệt (In-Browser RAG Context Extraction) với độ chính xác cao mà không cần dùng đến Vector Database trên đám mây."
    )
    add_numbered_item(
        doc, "2", "Ngân hàng hơn 500 câu hỏi trắc nghiệm",
        "Chúng em tự sưu tầm, chọn lọc và biên soạn lại hơn 500 câu hỏi bài tập từ các đề thi tốt nghiệp THPT tham khảo và đề thi học sinh giỏi. Nhóm em tự giải chi tiết từng câu, gán nhãn độ khó từ mức cơ bản đến nâng cao để đưa vào phần Đấu trường PvP cho các bạn học sinh trong trường thi đấu cọ xát kiến thức với nhau. Toàn bộ dữ liệu này được cấu trúc hóa dưới dạng tệp dữ liệu tĩnh (.json) và nhúng thẳng vào mã nguồn gói ứng dụng (Static In-Memory Bundling), cho phép hệ thống tra cứu siêu tốc offline với độ trễ tính bằng mili-giây và chi phí vận hành 0 đồng."
    )
    add_numbered_item(
        doc, "3", "Dữ liệu 144 phản ứng Hóa học",
        "Nhóm em đã tra cứu bảng tính tan, sổ tay Hóa học và sách giáo khoa để tự tổng hợp 144 trường hợp trộn hóa chất thường gặp. Với mỗi phản ứng, nhóm ghi rõ hiện tượng đổi màu, kết tủa hoặc sủi bọt khí và viết phương trình ion thu gọn để máy tính hiển thị mô phỏng giống y hệt khi làm thí nghiệm thật ở trường. Dữ liệu phản ứng cũng được cấu trúc hóa dưới dạng file tĩnh .json nhúng sẵn, giúp phòng thí nghiệm phản hồi tức thì mà không cần mạng Internet."
    )
    add_numbered_item(
        doc, "4", "Dữ liệu gần 400 câu giao tiếp của học sinh (392 mẫu intent)",
        "Để Trợ lý Nova nói chuyện tự nhiên và hiểu được cách xưng hô của học sinh, nhóm em đã thử nghiệm thu thập 392 câu nói, câu hỏi thường ngày của các bạn cùng lớp khi hỏi bài (kể cả các từ viết tắt hay cách nói quen thuộc của lứa tuổi học trò) để dạy cho AI hiểu được ý định và chuyển đúng trang mà bạn cần. Bộ dữ liệu mẫu này được đóng gói sẵn dưới dạng file .json phục vụ huấn luyện và dự đoán trực tiếp ngay trên máy của học sinh."
    )
    add_numbered_item(
        doc, "5", "Quần thể 20 AI Bot luyện tập",
        "Chúng em tự tạo 20 tài khoản Bot có mức điểm ELO từ thấp đến cao (từ 50 đến 130 điểm) để khi các bạn học sinh vào Đấu trường mà chưa có bạn nào ghép cặp thì hệ thống sẽ tự động đưa Bot ra thi đấu cùng. Toàn bộ hồ sơ, thông số hạt giống và chiến thuật của 20 Bot đều được lưu trữ dưới dạng file tĩnh .json, cho phép hệ thống khởi tạo tức thời và nâng cấp các thế hệ Bot hoàn toàn offline mà không tạo tải trọng lên máy chủ."
    )
    
    # IV. CÁC DỊCH VỤ ĐÁM MÂY VÀ WEB API HỖ TRỢ
    add_heading(doc, "IV. CÁC DỊCH VỤ ĐÁM MÂY VÀ WEB API HỖ TRỢ")
    add_paragraph(
        doc,
        "Trong quá trình xây dựng hệ thống web, nhóm em đã phân định rõ ràng giữa các Dịch vụ Đám mây (Cloud Services) phục vụ đồng bộ dữ liệu và các Tiêu chuẩn Web API của Trình duyệt (Browser Web APIs) để đảm bảo hệ thống vận hành trơn tru với chi phí 0 đồng:"
    )
    
    add_numbered_item(
        doc, "1", "Dịch vụ Đám mây Firebase của Google (Cloud Firestore & Authentication)",
        "Nhóm sử dụng gói Spark miễn phí (0 VNĐ) để quản lý tài khoản xác thực người dùng, đồng bộ kết quả thi đấu thời gian thực giữa hai người chơi trên Đấu trường và lưu trữ bảng xếp hạng ELO toàn trường."
    )
    add_numbered_item(
        doc, "2", "Cổng định tuyến Không máy chủ Cloudflare Workers (Serverless Edge Gateway)",
        "Nhóm em tự thiết lập một Worker đóng vai trò Reverse Proxy trung gian ở rìa mạng (Edge Network) nhằm che giấu và bảo mật mã khóa API của Google Gemini, đồng thời tích hợp cơ chế giới hạn tần suất (Rate Limiting) để bảo vệ tài nguyên miễn phí của dự án."
    )
    add_numbered_item(
        doc, "3", "Tiêu chuẩn Xử lý Đa luồng Web Worker (Off-Main-Thread Processing API)",
        "Khi học sinh đặt câu hỏi hoặc hệ thống tính toán ma trận TF-IDF / giải thuật di truyền, khối lượng tính toán khá lớn. Nhóm em đã chủ động tách toàn bộ các tác vụ nặng này sang một luồng ngầm độc lập (Background Worker Thread) thông qua Web Worker API. Nhờ đó, luồng giao diện chính (Main UI Thread) luôn được giải phóng, duy trì ổn định tốc độ hiển thị 60 FPS mượt mà, không bao giờ bị đơ hay giật lag kể cả trên các máy tính cấu hình yếu tại trường học."
    )
    add_numbered_item(
        doc, "4", "Tiêu chuẩn Nhận diện và Xử lý Âm thanh Web Speech API & Web Audio API",
        "Nhóm tận dụng công nghệ Web Speech API có sẵn của trình duyệt để học sinh có thể giao tiếp giọng nói tiếng Việt hai chiều với Trợ lý ảo Nova mà không tốn phí dịch vụ bên ngoài, kết hợp Web Audio API để xử lý hiệu ứng làm dịu nhạc nền (Audio Ducking) khi AI phản hồi."
    )
    add_numbered_item(
        doc, "5", "Tiêu chuẩn Cơ sở dữ liệu Cục bộ IndexedDB (Client-Side Storage API)",
        "Nhóm sử dụng IndexedDB để lưu trữ bền vững các tệp trọng số mô hình AI, bộ nhớ ngữ cảnh trò chuyện và dữ liệu thí nghiệm ngay tại trình duyệt máy học sinh. Điều này giúp loại bỏ giới hạn 5MB chật hẹp của LocalStorage, triệt tiêu nguy cơ lỗi tràn bộ nhớ (QuotaExceededError) và tăng tốc độ khởi động ứng dụng trong các lần truy cập tiếp theo."
    )
    
    # V. THƯ VIỆN MÃ NGUỒN MỞ
    add_heading(doc, "V. CÁC THƯ VIỆN MÃ NGUỒN MỞ MÀ NHÓM EM SỬ DỤNG")
    add_paragraph(
        doc,
        "Để giao diện trang web hiện đại, trực quan và dễ sử dụng cho các bạn học sinh, nhóm em đã chọn lọc sử dụng các thư viện mã nguồn mở miễn phí rất phổ biến trong cộng đồng lập trình hiện nay:"
    )
    
    add_numbered_item(
        doc, "1", "Next.js (phiên bản 16) và React (phiên bản 19)",
        "Đây là bộ khung chính (giấy phép mã nguồn mở MIT) mà nhóm em dùng để xây dựng toàn bộ trang web và kết nối các trang bài học với nhau."
    )
    add_numbered_item(
        doc, "2", "Thư viện Three.js và React Three Fiber",
        "Giúp nhóm dựng các mô hình 3D (giấy phép MIT) cho phòng thí nghiệm Vật lý như con lắc chuyển động hay đường đi khúc xạ của tia sáng laser qua các khối lăng kính."
    )
    add_numbered_item(
        doc, "3", "Thư viện KaTeX",
        "Giúp trang web hiển thị các công thức Toán học, công thức cấu tạo Hóa học và phương trình Vật lý (giấy phép MIT) rõ nét, chuẩn mực y hệt như in trong sách giáo khoa."
    )
    add_numbered_item(
        doc, "4", "Thư viện Tailwind CSS và Framer Motion",
        "Nhóm dùng để phối màu giao diện theo phong cách phòng nghiên cứu khoa học và làm các hiệu ứng mở cửa sổ, chuyển trang mượt mà (giấy phép MIT), tạo cảm hứng học tập cho học sinh."
    )
    add_numbered_item(
        doc, "5", "Thư viện Zustand",
        "Giúp nhóm quản lý trạng thái điểm số tạm thời của hai người chơi trong lúc đang thi đấu 1v1 trên Đấu trường (giấy phép MIT), giúp trận đấu diễn ra tức thì mà không làm nghẽn mạng."
    )
    add_numbered_item(
        doc, "6", "Thư viện Lucide React",
        "Cung cấp các biểu tượng đồ họa khoa học và giáo dục (giấy phép ISC) rất trực quan trên thanh điều hướng của trang web."
    )
    add_numbered_item(
        doc, "7", "Thư viện html2canvas-pro và SheetJS (XLSX)",
        "Giúp các bạn học sinh có thể tải ảnh chụp giấy chứng nhận thành tích sau mỗi tuần học và hỗ trợ nhóm nhập danh sách câu hỏi trắc nghiệm từ file Excel vào hệ thống."
    )
    
    # VI. LỜI CAM KẾT DANH DỰ
    add_heading(doc, "VI. LỜI CAM KẾT DANH DỰ CỦA NHÓM HỌC SINH")
    add_paragraph(
        doc,
        "Nhóm học sinh chúng em xin cam đoan trước Ban Bí thư Trung ương Đoàn và Ban Giám khảo Cuộc thi:"
    )
    add_paragraph(
        doc,
        "Thứ nhất, toàn bộ những nội dung kê khai ở trên là hoàn toàn trung thực và chính xác. Tất cả các ý tưởng sư phạm, cấu trúc bài toán, bộ ngân hàng câu hỏi, dữ liệu phản ứng hóa học và phần lớn mã nguồn giao diện tương tác đều do nhóm chúng em cùng thầy giáo hướng dẫn tự tay nghiên cứu, thiết kế và phát triển. Các công cụ AI chỉ đóng vai trò hỗ trợ kỹ thuật, thúc đẩy quá trình học hỏi và hoàn thiện sản phẩm của chúng em."
    )
    add_paragraph(
        doc,
        "Thứ hai, nhóm chúng em cam kết tuân thủ nghiêm túc các giấy phép mã nguồn mở (MIT, Apache 2.0, ISC), tuyệt đối không sao chép nguyên mẫu hay sử dụng mã nguồn vi phạm quyền sở hữu trí tuệ của bất kỳ tổ chức hay cá nhân nào. Nhóm xin hoàn toàn chịu trách nhiệm trước Ban tổ chức về tính trung thực của văn bản kê khai này."
    )
    
    # VII. KÝ TÊN XÁC NHẬN
    p_date = doc.add_paragraph()
    p_date.paragraph_format.space_before = Pt(14)
    p_date.paragraph_format.space_after = Pt(8)
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_d = p_date.add_run("Nghệ An, ngày 03 tháng 09 năm 2026")
    r_d.font.italic = True
    r_d.font.size = Pt(12)
    
    # Signatures as natural text blocks
    p_sig1 = doc.add_paragraph()
    p_sig1.paragraph_format.space_after = Pt(4)
    r_gv_t = p_sig1.add_run("Thầy giáo hướng dẫn:\n")
    r_gv_t.bold = True
    r_gv_t.font.color.rgb = RGBColor(0x00, 0x28, 0x55)
    r_gv_n = p_sig1.add_run("Thầy Nguyễn Văn Phúc (Giáo viên Trường THPT Quỳnh Lưu 4)")
    
    p_sig2 = doc.add_paragraph()
    p_sig2.paragraph_format.space_before = Pt(6)
    p_sig2.paragraph_format.space_after = Pt(2)
    r_tm_t = p_sig2.add_run("Đại diện nhóm học sinh tác giả:\n")
    r_tm_t.bold = True
    r_tm_t.font.color.rgb = RGBColor(0x00, 0x28, 0x55)
    r_tm_1 = p_sig2.add_run("Hồ Hoàng Anh (Học sinh lớp 12A, Trưởng nhóm)\n")
    r_tm_1.bold = True
    r_tm_2 = p_sig2.add_run("Các thành viên trong nhóm: Nguyễn Trung Hiếu (Lớp 12A), Đậu Việt Hưng (Lớp 12A)")
    r_tm_2.italic = True
    
    doc.save(output_path)
    print(f"✅ Đã tạo thành công file kê khai theo phong cách học sinh lớp 12: {output_path}")

if __name__ == "__main__":
    path1 = "e:/project/cuoc_thi_AI/Kê khai công cụ AI, dữ liệu, API, thư viện, mã nguồn mở.docx"
    path2 = "e:/project/cuoc_thi_AI/Ke_Khai_Cong_Cu_AI_Du_Lieu_API_Thu_Vien_Ma_Nguon_Mo.docx"
    path3 = "e:/project/cuoc_thi_AI/4_Ke_Khai_Cong_Cu_AI_Du_Lieu_API_Thu_Vien.docx"
    
    create_natural_student_doc(path1)
    create_natural_student_doc(path2)
    create_natural_student_doc(path3)
    print("HOÀN TẤT CẬP NHẬT TẤT CẢ CÁC FILE KÊ KHAI!")
