# -*- coding: utf-8 -*-
"""
Chuyển đổi LICHSU_CAULENH_VA_MINHCHUNG.md → Lich_Su_Cau_Lenh_Va_Hinh_Anh_Minh_Chung.docx
Sử dụng python-docx 1.2.0
"""

import re
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Paths ────────────────────────────────────────────────────────────────────
ROOT = Path(__file__).resolve().parent.parent
MD_FILE  = ROOT / "LICHSU_CAULENH_VA_MINHCHUNG.md"
OUT_FILE = ROOT / "Lich_Su_Cau_Lenh_Va_Hinh_Anh_Minh_Chung.docx"

# ── Helpers ──────────────────────────────────────────────────────────────────
def set_run_bold(run, bold: bool = True):
    run.bold = bold

def strip_md_inline(text: str) -> list:
    """
    Returns list of (text, bold) tuples by parsing **bold** markers.
    Does NOT recurse into code spans.
    """
    parts = []
    pattern = re.compile(r'\*\*(.+?)\*\*')
    last = 0
    for m in pattern.finditer(text):
        if m.start() > last:
            parts.append((text[last:m.start()], False))
        parts.append((m.group(1), True))
        last = m.end()
    if last < len(text):
        parts.append((text[last:], False))
    return parts if parts else [(text, False)]

def add_para_with_inline(doc, text: str, style_name: str = "Normal"):
    p = doc.add_paragraph(style=style_name)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    clean = text.lstrip(">").strip()
    # Remove leading markdown emphasis markers *...* that remain
    clean = re.sub(r'^\*(.+)\*$', r'\1', clean)
    for part_text, bold in strip_md_inline(clean):
        run = p.add_run(part_text)
        run.bold = bold
    return p

def add_heading(doc, text: str, level: int):
    # Strip leading # and emoji prefixes
    clean = re.sub(r'^#+\s*', '', text).strip()
    clean = re.sub(r'^📍\s*', '', clean).strip()
    h = doc.add_heading(clean, level=level)
    return h

def add_code_block(doc, lines: list):
    """Add a code / preformatted block with light-grey shading."""
    code_text = "\n".join(lines)
    p = doc.add_paragraph(style="Normal")
    run = p.add_run(code_text)
    run.font.name = "Courier New"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    # Light grey background via paragraph shading
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F0F0F0')
    pPr.append(shd)
    # Left indent
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), '360')
    pPr.append(ind)
    return p

def add_table_from_md(doc, header_row: list, body_rows: list):
    cols = len(header_row)
    table = doc.add_table(rows=1 + len(body_rows), cols=cols)
    table.style = "Table Grid"
    # Header
    hdr_cells = table.rows[0].cells
    for i, cell_text in enumerate(header_row):
        hdr_cells[i].text = cell_text.strip()
        for run in hdr_cells[i].paragraphs[0].runs:
            run.bold = True
        # Blue-grey fill for header
        tc = hdr_cells[i]._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), '2E4057')
        tcPr.append(shd)
        for para in hdr_cells[i].paragraphs:
            for run in para.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    # Body rows
    for r_idx, row in enumerate(body_rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, cell_text in enumerate(row):
            cells[c_idx].text = cell_text.strip()
    return table

def parse_md_table(lines: list):
    """Parse markdown table lines → (header_list, body_list_of_lists)"""
    if not lines:
        return [], []
    header = [c.strip() for c in lines[0].strip().strip('|').split('|')]
    body = []
    for line in lines[2:]:  # skip separator row
        if line.strip().startswith('|'):
            row = [c.strip() for c in line.strip().strip('|').split('|')]
            body.append(row)
    return header, body

# ── Main conversion ───────────────────────────────────────────────────────────
def convert():
    md_text = MD_FILE.read_text(encoding="utf-8")
    lines = md_text.splitlines()

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2)

    # Default Normal font
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)

    i = 0
    in_code = False
    code_lines = []
    table_lines = []
    in_table = False

    while i < len(lines):
        line = lines[i]

        # ── Code block ────────────────────────────────────────────────────
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
                i += 1
                continue
            else:
                in_code = False
                add_code_block(doc, code_lines)
                i += 1
                continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # ── Table ─────────────────────────────────────────────────────────
        if line.strip().startswith("|"):
            table_lines.append(line)
            i += 1
            continue
        else:
            if table_lines:
                header, body = parse_md_table(table_lines)
                if header:
                    add_table_from_md(doc, header, body)
                    doc.add_paragraph()  # spacing after table
                table_lines = []

        # ── Horizontal rule ───────────────────────────────────────────────
        if re.match(r'^-{3,}$', line.strip()):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), '999999')
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue

        # ── Headings ──────────────────────────────────────────────────────
        if line.startswith("# ") and not line.startswith("## "):
            add_heading(doc, line, level=1)
            i += 1
            continue
        if line.startswith("## ") and not line.startswith("### "):
            add_heading(doc, line, level=2)
            i += 1
            continue
        if line.startswith("### "):
            add_heading(doc, line, level=3)
            i += 1
            continue
        if line.startswith("#### "):
            add_heading(doc, line, level=4)
            i += 1
            continue

        # ── Blockquote ────────────────────────────────────────────────────
        if line.startswith(">"):
            clean = line.lstrip(">").strip()
            clean = re.sub(r'^\*(.+)\*$', r'\1', clean)
            p = doc.add_paragraph(style="Normal")
            p.paragraph_format.left_indent = Inches(0.4)
            for part_text, bold in strip_md_inline(clean):
                run = p.add_run(part_text)
                run.bold = bold
                run.italic = True
                run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            i += 1
            continue

        # ── Bullet list ───────────────────────────────────────────────────
        if re.match(r'^[-*]\s+', line.strip()):
            clean = re.sub(r'^[-*]\s+', '', line.strip())
            p = doc.add_paragraph(style="List Bullet")
            for part_text, bold in strip_md_inline(clean):
                run = p.add_run(part_text)
                run.bold = bold
            i += 1
            continue

        # ── Numbered list ─────────────────────────────────────────────────
        if re.match(r'^\d+\.\s+', line.strip()):
            clean = re.sub(r'^\d+\.\s+', '', line.strip())
            p = doc.add_paragraph(style="List Number")
            for part_text, bold in strip_md_inline(clean):
                run = p.add_run(part_text)
                run.bold = bold
            i += 1
            continue

        # ── Empty line ────────────────────────────────────────────────────
        if line.strip() == "":
            i += 1
            continue

        # ── Normal paragraph ──────────────────────────────────────────────
        add_para_with_inline(doc, line, "Normal")
        i += 1

    # flush remaining table
    if table_lines:
        header, body = parse_md_table(table_lines)
        if header:
            add_table_from_md(doc, header, body)

    doc.save(str(OUT_FILE))
    print(f"Saved: {OUT_FILE}  ({OUT_FILE.stat().st_size:,} bytes)")

if __name__ == "__main__":
    convert()
