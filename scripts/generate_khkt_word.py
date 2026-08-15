import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_khkt_student_report():
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(17)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(16)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(11)
        run.font.italic = True
        run.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x03, 0x69, 0xA1)

    def add_box(title, content, bg_color="F1F5F9", border_color="0284C7"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        
        tcPr = cell._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
        tcPr.append(shd)
        
        tcBorders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        tcPr.append(tcBorders)

        tcMar = parse_xml(f'''
            <w:tcMar {nsdecls("w")}>
                <w:top w:w="120" w:type="dxa"/>
                <w:left w:w="180" w:type="dxa"/>
                <w:bottom w:w="120" w:type="dxa"/>
                <w:right w:w="180" w:type="dxa"/>
            </w:tcMar>
        ''')
        tcPr.append(tcMar)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(4)
        run_title = p.add_run(title)
        run_title.font.name = 'Arial'
        run_title.font.size = Pt(10.5)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.15
        run_content = p2.add_run(content)
        run_content.font.name = 'Arial'
        run_content.font.size = Pt(10)
        run_content.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # HEADER
    add_title("KỊCH BẢN NÓI KHKT DÀNH CHO HỌC SINH (NGÔN NGỮ TỰ NHIÊN, DỄ HIỂU)")
    add_subtitle("Dự án: Hệ thống Mô phỏng và Đấu trường Tri thức STEM\nGiúp học sinh tự tin bảo vệ mà không bị ngập trong thuật toán phức tạp")

    # STEP 1
    add_h1("BƯỚC 1: CÁCH HỌC SINH GIỚI THIỆU TỔNG QUAN DỰ ÁN (30 GIÂY)")
    p = doc.add_paragraph()
    p.add_run("Dạ thưa Ban giám khảo, dự án của tụi em là một phần mềm mô phỏng giúp học sinh học Vật lý và Toán 3D trực quan ngay trên trang web:\n")
    p.add_run("• Biến công thức khô khan thành hình ảnh tương tác: ").bold = True
    p.add_run("Khi kéo các thanh trượt thay đổi khối lượng, góc quay hay tần số, hình ảnh 3D và đồ thị trên màn hình sẽ biến đổi theo ngay lập tức.\n")
    p.add_run("• Chạy mượt trên mọi máy tính: ").bold = True
    p.add_run("Tụi em lập trình để trang web tự động chia việc cho card màn hình xử lý hình ảnh, nên máy tính cũ ở phòng lab của trường hay điện thoại cũng chạy rất mượt.\n")
    p.add_run("• Giải quyết bài toán thực tế: ").bold = True
    p.add_run("Giúp học sinh học được những bài thí nghiệm đắt tiền hoặc nguy hiểm mà phòng thí nghiệm ở trường chưa có điều kiện làm.")

    # STEP 2 & 3
    add_h1("BƯỚC 2 & BƯỚC 3: CÂU HỎI VÀ LỜI ĐÁP THEO GIỌNG VĂN HỌC SINH (DỄ HIỂU, ĐỦ Ý)")

    add_h2("NHÓM 1: CÂU HỎI VỀ CÁCH HOẠT ĐỘNG VÀ TÍNH ĐÚNG ĐẮN VẬT LÝ")

    # Q1
    add_box(
        "Câu 1 (Về tính chính xác vật lý):\nGiám khảo: Làm sao em chắc chắn phần mềm chạy đúng với thực tế? Việc dùng công thức đơn giản hay bỏ qua lực cản không khí có làm sai lệch kết quả không?",
        "LỜI ĐÁP DỄ HIỂU CỦA HỌC SINH:\n"
        "1. Dạ, em quy đổi 1 mét thực tế thành một số điểm ảnh cố định trong màn hình 3D, và dùng đồng hồ thực tế của máy tính để tính thời gian, nên xe chạy hay con lắc lắc nhanh chậm đều đúng tỉ lệ thực tế.\n"
        "2. Về sai số: Với bài con lắc đơn, công thức trong SGK lớp 12 tụi em dùng chỉ đúng khi kéo góc nhỏ dưới 10 độ. Nếu kéo tới 30 độ thì năng lượng thực tế sẽ bị lệch một chút.\n"
        "3. Hướng khắc phục: Em chọn cách làm này ở bản đầu để máy tính yếu của học sinh cũng chạy mượt. Sau này em sẽ lập trình thêm công thức tính thực tế hơn có tính cả lực cản không khí ạ.",
        "F1F5F9", "0284C7"
    )

    # Q2
    add_box(
        "Câu 2 (Về độ mượt khi vừa vẽ 3D vừa hiện công thức):\nGiám khảo: Vừa quay hình 3D, vừa vẽ đồ thị lại vừa hiện công thức toán có làm máy tính ở trường bị giật lag không? Em làm sao cho nó mượt?",
        "LỜI ĐÁP DỄ HIỂU CỦA HỌC SINH:\n"
        "1. Dạ, tụi em tách riêng công việc: Màn hình 3D thì cứ tự quay mượt mà 60 hình/giây, không bị ảnh hưởng bởi các chữ viết xung quanh.\n"
        "2. Với các công thức toán: Khi nào học sinh buông tay khỏi thanh trượt thì phần mềm mới hiện con số cố định ra, chứ không bắt máy tính vẽ lại chữ liên tục từng miligiây.\n"
        "3. Tụi em cũng đẩy hết việc tính toán màu sắc đồ thị sang cho Card đồ họa của máy tính làm, nên CPU không bị quá tải.",
        "F1F5F9", "0284C7"
    )

    # Q3
    add_box(
        "Câu 3 (Về cách cắt hình 3D lấy mặt cắt):\nGiám khảo: Trong bài hình học 3D, em làm thế nào để phần mềm tự tìm và vẽ ra cái mặt cắt (thiết diện) khi cắt một khối hình bằng một mặt phẳng?",
        "LỜI ĐÁP DỄ HIỂU CỦA HỌC SINH:\n"
        "1. Dạ, em cho phần mềm dò xem cái mặt phẳng cắt đó đi qua những cạnh nào của khối 3D để tìm các điểm giao.\n"
        "2. Sau đó, phần mềm tự động lọc bỏ các điểm trùng nhau và sắp xếp các điểm đó theo đúng một vòng xung quanh tâm để nối lại thành hình mặt cắt kín.\n"
        "3. Nhờ cách tính này, máy tính học sinh tính xong mặt cắt trong chưa tới 1 phần ngàn giây mà không cần gửi lên mạng.",
        "F1F5F9", "0284C7"
    )

    # Q4
    add_box(
        "Câu 4 (Về thi đấu 1v1 trực tuyến):\nGiám khảo: Khi 2 bạn ở 2 nơi cùng thi đấu 1v1, bạn dùng 3G mạng yếu hơn có bị thiệt thòi về thời gian so với bạn dùng Wifi không?",
        "LỜI ĐÁP DỄ HIỂU CỦA HỌC SINH:\n"
        "1. Dạ không ạ. Thời gian bắt đầu và nộp bài được tính bằng giờ đồng hồ trên máy chủ (Server) chứ không dùng giờ trên điện thoại học sinh nên không gian lận được.\n"
        "2. Trước khi đấu, phần mềm sẽ tự đo tốc độ mạng của 2 bạn. Bạn nào mạng 3G chậm hơn bao nhiêu milligiây thì hệ thống sẽ tự động cộng thêm bấy nhiêu thời gian bù trễ cho bạn đó.\n"
        "3. Khi cả 2 bạn cùng bấm đáp án một lúc, máy chủ sẽ tự xếp hàng phân xử trước sau rất công bằng.",
        "F1F5F9", "0284C7"
    )

    # Q5
    add_box(
        "Câu 5 (Về đo sóng âm qua Micro):\nGiám khảo: Em tính tần số sóng âm từ Micro thế nào? Tại sao đo tiếng nói từ Micro lại dễ bị nhảy số linh tinh hơn âm thanh từ máy phát?",
        "LỜI ĐÁP DỄ HIỂU CỦA HỌC SINH:\n"
        "1. Dạ, phần mềm thu tiếng từ Micro rồi đo khoảng cách giữa các đỉnh sóng để tính ra chu kỳ và tần số âm thanh.\n"
        "2. Âm từ máy phát là tiếng tít thuần khiết (sóng sin mịn) nên đếm rất dễ. Còn tiếng nói người qua Micro là âm phức tạp có lẫn tạp âm xung quanh, làm đồ thị bị nhấp nhô răng cưa nên máy hay đếm nhầm.\n"
        "3. Em đã làm thêm một bộ lọc tiếng ồn để lọc bớt nhiễu và lấy giá trị trung bình qua nhiều lần đo để con số ổn định.",
        "F1F5F9", "0284C7"
    )

    # GROUP 2
    add_h2("NHÓM 2: CÂU HỎI THỬ THÁCH & XÁC THỰC THỰC LỰC HỌC SINH")

    # Q6
    add_box(
        "Câu 6 (Tại sao không dùng phần mềm làm game Unity):\nGiám khảo: Tại sao em không dùng mấy phần mềm làm game có sẵn như Unity cho nhanh mà phải tự lập trình trên trang web từ đầu?",
        "LỜI ĐÁP DỄ HIỂU CỦA HỌC SINH:\n"
        "1. Dạ, nếu làm bằng Unity thì file rất nặng (hàng trăm MB), học sinh muốn dùng phải tải về cài đặt rất phức tạp, điện thoại yếu chạy không nổi.\n"
        "2. Làm trên trang web thì học sinh chỉ cần bấm vào link là chạy ngay trong 2 giây, dùng được trên cả điện thoại và máy tính ở trường.\n"
        "3. Tự lập trình giúp em hiểu rõ tận gốc toán học và lý thuyết đằng sau chứ không chỉ đi bấm nút có sẵn.",
        "FEF2F2", "EF4444"
    )

    # Q7
    add_box(
        "Câu 7 (Về âm thanh 3D khi lại gần ra xa):\nGiám khảo: Âm thanh trong bài của em nghe cố định hay có tự to nhỏ/thay đổi khi nguồn âm di chuyển lại gần hoặc ra xa trong không gian 3D?",
        "LỜI ĐÁP DỄ HIỂU CỦA HỌC SINH:\n"
        "1. Dạ ở bản hiện tại, phần mềm hoạt động như một máy đo tần số 2D nên âm thanh phát ra cố định theo số mình chỉnh.\n"
        "2. Em đã nghiên cứu xong cách làm âm thanh 3D, ở bản tiếp theo khi nguồn âm chạy lại gần tiếng sẽ cao và to hơn, đi ra xa tiếng sẽ trầm và nhỏ dần đúng thực tế ạ.",
        "FEF2F2", "EF4444"
    )

    # Q8
    add_box(
        "Câu 8 (Khi kéo liên tục thanh trượt):\nGiám khảo: Khi kéo thanh trượt chỉnh khối lượng hay độ cứng, nếu hình vẽ cứ bắt trang web vẽ lại liên tục thì màn hình có bị khựng không?",
        "LỜI ĐÁP DỄ HIỂU CỦA HỌC SINH:\n"
        "1. Dạ có, nếu cứ kéo 1 miligiây mà bắt trang web tính lại hết thì màn hình sẽ bị đơ ngay.\n"
        "2. Nên tụi em xử lý là: Khi đang giữ chuột kéo thì hình 3D cứ tự quay mượt mà theo tay mình, chừng nào mình buông tay ra thì con số mới hiện cố định lên màn hình.\n"
        "3. Nhờ vậy mà thao tác kéo thả rất mượt mà không bị giật lag.",
        "FEF2F2", "EF4444"
    )

    # Q9
    add_box(
        "Câu 9 (Về xoay gương tia Laser góc lẻ):\nGiám khảo: Trong bài tia Laser, nếu em xoay gương một góc lẻ như 33.5 độ thì tia phản xạ có chiếu đúng không hay chỉ xoay được các góc chẵn?",
        "LỜI ĐÁP DỄ HIỂU CỦA HỌC SINH:\n"
        "1. Dạ ở bản v1, vì em cho gương và laser di chuyển theo các ô vuông cố định nên phần mềm chỉ mới hỗ trợ các góc chẵn như 45 hay 90 độ.\n"
        "2. Trong bản cập nhật em vừa làm xong, em đã áp dụng công thức vector đường thẳng nên gương xoay góc lẻ bao nhiêu độ như 33.5 độ thì tia laser vẫn phản xạ chính xác theo pháp tuyến ạ.",
        "FEF2F2", "EF4444"
    )

    # Q10
    add_box(
        "Câu 10 (Điểm tiếc nuối nhất và hướng phát triển):\nGiám khảo: Nếu chọn MỘT điều em chưa hài lòng nhất ở sản phẩm này, đó là gì và em sẽ sửa nó thế nào trong 6 tháng tới?",
        "LỜI ĐÁP DỄ HIỂU CỦA HỌC SINH:\n"
        "1. Dạ, điều em tiếc nuối nhất là hiện tại máy tính vừa phải lo vẽ hình 3D vừa phải tính toán công thức vật lý trên cùng một luồng nên một số bài phức tạp em phải dùng công thức xấp xỉ đơn giản.\n"
        "2. Trong 6 tháng tới, em sẽ tách phần tính toán vật lý sang một luồng chạy ngầm riêng, giúp phần mềm vừa tính được va chạm phức tạp có lực cản, vừa giữ được màn hình mượt mà 60 hình/giây.",
        "FEF2F2", "EF4444"
    )

    filename = "e:/Downloads/STT2026/Bao_cao_Phan_bien_KHKT_Duy_an_STEM_2026_HocSinh.docx"
    doc.save(filename)
    print(f"File updated successfully: {filename}")

if __name__ == "__main__":
    create_khkt_student_report()
